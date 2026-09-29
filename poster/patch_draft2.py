"""
Apply the requested changes on top of Sasha's own draft (draft1.pptx), keeping
every word of his edits except the ones he asked to change.

  - term cards say which names are rivers and which is a city
  - standfirst opens with the full downstream path, naming each place in the
    order the water meets it (his second half is kept exactly as he wrote it)
  - long profile and the two sections swapped for regenerated images
  - background rectangle put back at y = 0 (it had drifted 3.5 mm down,
    leaving a pale strip under the hazard rule)
"""
import copy
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Mm, Pt

SRC, OUT, ASSETS = sys.argv[1], sys.argv[2], sys.argv[3]

INK = RGBColor(0x14, 0x18, 0x1C)
GREY = RGBColor(0x6A, 0x65, 0x50)

prs = Presentation(SRC)
slide = prs.slides[0]
by_name = {sh.name: sh for sh in slide.shapes}


def must(name):
    if name not in by_name:
        sys.exit("missing shape: " + name)
    return by_name[name]


# ---------------------------------------------------------- 1. background
bg = must("BACKGROUND — hazard yellow")
bg.top = Emu(0)
bg.height = prs.slide_height

# ------------------------------------------- 2. rivers vs city on the cards
KIND = {"TYUMEN": "city", "TURA": "river", "TOBOL": "river",
        "IRTYSH": "river", "OB": "river"}
for key, kind in KIND.items():
    tb = must("term — %s" % key)
    p0 = tb.text_frame.paragraphs[0]
    title = p0.runs[0]
    if title.text.strip().upper() != key:
        sys.exit("unexpected card title: %r" % title.text)
    # a second run, same size, grey and not bold
    new_r = copy.deepcopy(title._r)
    p0._p.append(new_r)
    r = p0.runs[-1]
    r.text = "  ·  " + kind
    r.font.bold = False
    r.font.color.rgb = GREY

# ------------------------------------------------------ 3. the standfirst
sf = must("standfirst").text_frame
para = sf.paragraphs[0]
runs = para.runs
full = "".join(r.text for r in runs)
# keep everything from "It does four-fifths" onwards exactly as written
anchor = "It does four-fifths"
k = full.find(anchor)
if k < 0:
    sys.exit("could not find the part of the standfirst to keep")

# locate the run in which the kept text starts, and split it there
acc = 0
for idx, r in enumerate(runs):
    if acc + len(r.text) > k:
        split_at = k - acc
        break
    acc += len(r.text)
keep_first = runs[idx]
head_text = keep_first.text[:split_at]
keep_first.text = keep_first.text[split_at:]
for r in runs[:idx]:
    para._p.remove(r._r)

template = keep_first._r
size = keep_first.font.size or Pt(11.3)


def insert_before(anchor_r, text, bold):
    new = copy.deepcopy(template)
    anchor_r.addprevious(new)
    from pptx.text.text import _Run
    run = _Run(new, para)
    run.text = text
    run.font.bold = bold
    run.font.size = size
    run.font.color.rgb = INK
    return new


OPENING = [
    ("The ", False), ("Tura", True),
    (" rises at 370 m in the Ural mountains and flows east past ", False),
    ("Tyumen", True), (", my home city, into the ", False), ("Tobol", True),
    (". The Tobol flows into the ", False), ("Irtysh", True),
    (", and the Irtysh into the ", False), ("Ob", True),
    (", which carries all of it north to the Kara Sea in the Arctic — four "
     "rivers, one journey from source to sea. ", False),
]
for text, bold in OPENING:
    insert_before(template, text, bold)
keep_first.font.bold = True           # "It does four-fifths ..." stays bold

# the box needs one more line
sf_shape = must("standfirst")
sf_shape.height = Mm(21.5)

# ------------------------------------------------------ 4. swap the images
for name, fname in [
        ("LONG PROFILE — Tura to Kara Sea (measured)", "long_profile.png"),
        ("SECTION — Tyumen, the Tura", "section_tyumen.png"),
        ("SECTION — Shrewsbury, the Severn", "section_shrewsbury.png")]:
    pic = must(name)
    rId = pic._element.blipFill.blip.rEmbed
    part = slide.part.related_part(rId)
    part._blob = open("%s/%s" % (ASSETS, fname), "rb").read()

prs.save(OUT)
print("saved", OUT)
print("standfirst now:", "".join(r.text for r in sf.paragraphs[0].runs))
